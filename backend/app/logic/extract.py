"""资料解析：把上传的文件切成带定位信息的“提取单元”（同学文档 DC-01、DC-03）。

定位格式：
- PDF 物理页        pdf:p005             第 5 页
- PPTX 幻灯片       pptx:s012            第 12 张幻灯片
- DOCX 段落         docx:para0012        第 12 段
- DOCX 表格行       docx:table002:r003   表 2 第 3 行
- 文本 / 代码行号    text:L0001-L0040     第 1–40 行（代码为 cpp: 或 code:）
- 图片              image:file           整张图片（装了识别组件时自动识别文字）
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

from . import ocr
from .text import normalize_extracted

TEXT_LINES_PER_UNIT = 40

FORMATS: dict[str, str] = {
    ".pdf": "pdf",
    ".pptx": "pptx",
    ".docx": "docx",
    ".txt": "text",
    ".md": "md",
    ".markdown": "md",
    ".cpp": "code",
    ".cc": "code",
    ".c": "code",
    ".h": "code",
    ".hpp": "code",
    ".py": "code",
    ".java": "code",
    ".png": "image",
    ".jpg": "image",
    ".jpeg": "image",
}

C_FAMILY = {".cpp", ".cc", ".c", ".h", ".hpp"}

CONVERT_HINTS = {
    ".ppt": "旧版 .ppt 请先在 PowerPoint 或 Keynote 中另存为 .pptx 再上传",
    ".doc": "旧版 .doc 请先在 Word 或 Pages 中另存为 .docx 再上传",
}

MAGIC = {
    "pdf": [b"%PDF"],
    "pptx": [b"PK\x03\x04"],
    "docx": [b"PK\x03\x04"],
    "image": [b"\x89PNG", b"\xff\xd8\xff"],
}

WARNING_TEXT = {
    "no_text": "这一页没有可提取的文字，可能是扫描图片",
    "extract_error": "这一部分解析出错，已跳过",
    "ocr_unavailable": "没有安装文字识别组件，图片里的文字需要手动补充",
    "ocr": "文字由图片识别得到，可能有错字，建议对照原图校对",
    "user_edited": "文字由你手动补充或修改",
    "blank": "内容为空",
}


@dataclass
class Unit:
    kind: str
    ordinal: int
    locator: str
    label: str
    text: str
    warnings: list[str] = field(default_factory=list)


@dataclass
class ExtractResult:
    status: str  # ready / partial / failed
    units: list[Unit]
    error_code: str | None = None
    error_message: str | None = None


def detect_format(filename: str) -> tuple[str | None, str, str | None]:
    """返回 (格式, 扩展名, 不支持时的原因)。"""
    ext = Path(filename).suffix.lower()
    if ext in FORMATS:
        return FORMATS[ext], ext, None
    if ext in CONVERT_HINTS:
        return None, ext, CONVERT_HINTS[ext]
    return None, ext, f"不支持 {ext or '无扩展名'} 文件，可上传 PDF、PPTX、DOCX、TXT、MD、代码或图片"


def check_content_type(fmt: str, head: bytes) -> bool:
    """按文件开头的字节核对内容类型，防止改了扩展名的文件进入解析。"""
    if fmt in MAGIC:
        return any(head.startswith(sig) for sig in MAGIC[fmt])
    # 文本类：前 8 KB 不应含有 NUL 字节
    return b"\x00" not in head[:8192]


def extract(path: Path, fmt: str, ext: str) -> ExtractResult:
    try:
        if fmt == "pdf":
            return _extract_pdf(path)
        if fmt == "pptx":
            return _extract_pptx(path)
        if fmt == "docx":
            return _extract_docx(path)
        if fmt in ("text", "md", "code"):
            return _extract_text(path, fmt, ext)
        if fmt == "image":
            return _extract_image(path)
    except Exception as exc:  # 解析库抛出的各种异常统一归为“文件损坏”
        return ExtractResult("failed", [], "corrupted", f"文件无法解析，可能已损坏（{type(exc).__name__}）")
    return ExtractResult("failed", [], "unsupported", "不支持的格式")


def _finish(units: list[Unit], kind_label: str) -> ExtractResult:
    if not units:
        return ExtractResult("partial", [], "empty", "没有提取到任何文字，可以手动补充")
    empty = [u for u in units if not u.text.strip()]
    if len(empty) == len(units):
        return ExtractResult(
            "partial", units, "no_text",
            "没有提取到文字，可能是扫描件；可以在对应位置手动补充文字",
        )
    if empty:
        return ExtractResult("partial", units, "partial_text", f"有 {len(empty)} {kind_label}没有可提取的文字")
    return ExtractResult("ready", units)


def _extract_pdf(path: Path) -> ExtractResult:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    if reader.is_encrypted:
        try:
            ok = reader.decrypt("")
        except Exception:
            ok = 0
        if not ok:
            return ExtractResult("failed", [], "encrypted", "这个 PDF 需要密码才能读取，请先解除密码再上传")
    units: list[Unit] = []
    for index, page in enumerate(reader.pages, start=1):
        warnings: list[str] = []
        try:
            text = normalize_extracted(page.extract_text() or "")
        except Exception:
            text = ""
            warnings.append("extract_error")
        if not text and not warnings:
            warnings.append("no_text")
        units.append(Unit("page", index, f"pdf:p{index:03d}", f"第 {index} 页", text, warnings))
    # 扫描版 PDF：没有文字层的页用 OCR 识别
    blank = [u.ordinal for u in units if not u.text.strip()]
    if blank and ocr.available():
        found = ocr.ocr_pdf_pages(path, blank)
        for u in units:
            text = normalize_extracted(found.get(u.ordinal, ""))
            if text:
                u.text = text
                u.warnings = ["ocr"]
    return _finish(units, "页")


def _extract_image(path: Path) -> ExtractResult:
    if not ocr.available():
        unit = Unit("image", 1, "image:file", "整张图片", "", ["ocr_unavailable"])
        return ExtractResult("partial", [unit], "ocr_unavailable", WARNING_TEXT["ocr_unavailable"])
    text = normalize_extracted(ocr.ocr_image_file(path))
    if not text:
        unit = Unit("image", 1, "image:file", "整张图片", "", ["no_text"])
        return ExtractResult("partial", [unit], "no_text", "图片里没有识别到文字，可以手动补充")
    return ExtractResult("ready", [Unit("image", 1, "image:file", "整张图片", text, ["ocr"])])


def _shape_texts(shape, out: list[str]) -> None:  # type: ignore[no-untyped-def]
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    try:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            for child in shape.shapes:
                _shape_texts(child, out)
            return
    except Exception:
        pass
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        for paragraph in shape.text_frame.paragraphs:
            text = paragraph.text.strip()
            if text:
                out.append(text)
    if getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                out.append(" | ".join(cells))


def _extract_pptx(path: Path) -> ExtractResult:
    from pptx import Presentation

    prs = Presentation(str(path))
    units: list[Unit] = []
    for index, slide in enumerate(prs.slides, start=1):
        parts: list[str] = []
        warnings: list[str] = []
        for shape in slide.shapes:
            try:
                _shape_texts(shape, parts)
            except Exception:
                warnings.append("extract_error")
        try:
            if slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame.text.strip()
                if notes:
                    parts.append(f"（备注）{notes}")
        except Exception:
            pass
        text = normalize_extracted("\n".join(parts))
        if not text and "extract_error" not in warnings:
            warnings.append("no_text")
        units.append(Unit("slide", index, f"pptx:s{index:03d}", f"第 {index} 张幻灯片", text, sorted(set(warnings))))
    return _finish(units, "张幻灯片")


def _extract_docx(path: Path) -> ExtractResult:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = docx.Document(str(path))
    units: list[Unit] = []
    para_index = 0
    table_index = 0
    ordinal = 0
    for block in document.iter_inner_content():
        if isinstance(block, Paragraph):
            para_index += 1
            text = normalize_extracted(block.text)
            if not text:
                continue
            style = (block.style.name if block.style is not None else "") or ""
            if style.lower().startswith("heading") or style.startswith("标题"):
                text = f"【{text}】"
            ordinal += 1
            units.append(Unit("paragraph", ordinal, f"docx:para{para_index:04d}", f"第 {para_index} 段", text))
        elif isinstance(block, Table):
            table_index += 1
            for row_index, row in enumerate(block.rows, start=1):
                seen: list[object] = []
                cells: list[str] = []
                for cell in row.cells:
                    if any(cell._tc is s for s in seen):  # 合并单元格会重复出现
                        continue
                    seen.append(cell._tc)
                    cells.append(normalize_extracted(cell.text).replace("\n", " "))
                text = " | ".join(cells).strip(" |")
                if not text:
                    continue
                ordinal += 1
                units.append(Unit(
                    "row", ordinal, f"docx:table{table_index:03d}:r{row_index:03d}",
                    f"表 {table_index} 第 {row_index} 行", text,
                ))
    return _finish(units, "段")


def _extract_text(path: Path, fmt: str, ext: str) -> ExtractResult:
    raw = path.read_bytes()
    try:
        content = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        return ExtractResult("failed", [], "encoding", "文件不是 UTF-8 编码，请先转换为 UTF-8 再上传")
    content = content.replace("\r\n", "\n").replace("\r", "\n")
    lines = content.split("\n")
    if lines and lines[-1] == "":
        lines.pop()
    prefix = "cpp" if ext in C_FAMILY else ("code" if fmt == "code" else "text")
    units: list[Unit] = []
    for start in range(0, len(lines), TEXT_LINES_PER_UNIT):
        chunk = lines[start:start + TEXT_LINES_PER_UNIT]
        first, last = start + 1, start + len(chunk)
        text = "\n".join(chunk).rstrip()
        warnings = [] if text.strip() else ["blank"]
        units.append(Unit(
            "lines", len(units) + 1, f"{prefix}:L{first:04d}-L{last:04d}",
            f"第 {first}–{last} 行", text, warnings,
        ))
    if not any(u.text.strip() for u in units):
        return ExtractResult("partial", units, "blank", "文件内容为空白，需要补充内容")
    return ExtractResult("ready", units)
